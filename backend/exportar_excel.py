import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime


def gerar_planilha_notas(
    dados_notas,
    caminho_saida="relatorio_notas_fiscais.xlsx",
    honorario_por_medico=0,
    darf_gps_por_medico=0
):
    """
    Gera um arquivo Excel completo contendo:

    1. Planilha Geral
    2. Uma aba para cada médico

    Os cálculos das abas dos médicos seguem a mesma regra
    utilizada na página individual do médico no Streamlit.
    """

    wb = Workbook()

    # ========================================================
    # ESTILOS
    # ========================================================

    fonte_titulo = Font(
        bold=True,
        size=14
    )

    fonte_cabecalho = Font(
        bold=True
    )

    fonte_total = Font(
        bold=True
    )

    alinhamento_centro = Alignment(
        horizontal="center",
        vertical="center"
    )

    alinhamento_esquerda = Alignment(
        horizontal="left",
        vertical="center"
    )

    borda = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    formato_moeda = 'R$ #,##0.00'

    # ========================================================
    # MÉDICOS ENCONTRADOS
    # ========================================================

    medicos = sorted(
        set(
            nota.get("Medico", "").strip()
            for nota in dados_notas
            if nota.get("Medico", "").strip()
        )
    )

    quantidade_medicos = len(medicos)

    # ========================================================
    # ABA GERAL
    # ========================================================

    ws = wb.active
    ws.title = "Planilha Geral"

    ws["A1"] = "PLANILHA DE FATURAMENTO"
    ws["A1"].font = fonte_titulo

    # ========================================================
    # CABEÇALHOS DAS NOTAS
    # ========================================================

    cabecalhos = [
        "DATA",
        "CLIENTE",
        "N.FISCAL",
        "VALOR",
        "IR",
        "INSS",
        "PIS/COFINS/CSLL",
        "VALOR LÍQUIDO",
        "MÉDICO"
    ]

    linha = 3

    for coluna, cabecalho in enumerate(
        cabecalhos,
        start=1
    ):

        celula = ws.cell(
            row=linha,
            column=coluna,
            value=cabecalho
        )

        celula.font = fonte_cabecalho
        celula.alignment = alinhamento_centro
        celula.border = borda

    linha += 1

    linha_inicio_notas = linha

    # ========================================================
    # NOTAS FISCAIS
    # ========================================================

    for nota in dados_notas:

        data = nota.get(
            "DataEmissao",
            ""
        )

        if isinstance(data, str):
            try:
                data = datetime.strptime(
                    data,
                    "%d/%m/%Y %H:%M:%S"
                )
            except ValueError:
                pass

        cliente = nota.get(
            "TomadorServico",
            ""
        )

        numero_nf = nota.get(
            "NumeroNF",
            ""
        )

        valor = float(
            nota.get("ValorTotal", 0)
        )

        ir = float(
            nota.get("ValorIr", 0)
        )

        inss = float(
            nota.get("ValorInss", 0)
        )

        situacao_nota = str(
            nota.get("SituacaoNota", "")
        ).strip().upper()

        valor_pcc_retido = nota.get(
            "ValorPccRetido"
        )

        csll = float(
            nota.get("ValorCsll", 0)
        )

        if situacao_nota == "C":
            pis_cofins_csll = 0

        elif (
            valor_pcc_retido is not None
            and valor_pcc_retido != ""
        ):
            pis_cofins_csll = float(
                valor_pcc_retido
            )

        elif csll > 0:
            valor_total = float(
                nota.get("ValorTotal", 0)
            )

            pis_cofins_csll = round(
                valor_total * 0.0465,
                2
            )

        else:
            pis_cofins_csll = 0

        medico = nota.get(
            "Medico",
            ""
        ).strip()

        situacao_nota = nota.get(
            "SituacaoNota",
            ""
        ).strip().upper()

        if situacao_nota == "C":
            valor = 0
            ir = 0
            inss = 0
            pis_cofins_csll = 0

        ws.cell(
            row=linha,
            column=1,
            value=data
        )
        ws.cell(
            row=linha,
            column=1
        ).number_format = "dd/mm/yyyy"

        ws.cell(
            row=linha,
            column=2,
            value=cliente
        )

        ws.cell(
            row=linha,
            column=3,
            value=numero_nf
        )

        valores = [
            valor,
            ir,
            inss,
            pis_cofins_csll
        ]

        for coluna, valor_celula in enumerate(
            valores,
            start=4
        ):

            celula = ws.cell(
                row=linha,
                column=coluna,
                value=valor_celula
            )

            celula.number_format = formato_moeda

        # VALOR LÍQUIDO
        celula_liquido = ws.cell(
            row=linha,
            column=8,
            value=(
                f"=D{linha}-E{linha}"
                f"-F{linha}-G{linha}"
            )
        )

        celula_liquido.number_format = formato_moeda

        ws.cell(
            row=linha,
            column=9,
            value=medico
        )

        for coluna in range(1, 10):

            ws.cell(
                row=linha,
                column=coluna
            ).border = borda

        linha += 1

    linha_fim_notas = linha - 1
    # ========================================================
    # TOTAL DAS NOTAS
    # ========================================================

    linha_total_notas = linha

    ws.cell(
        row=linha_total_notas,
        column=3,
        value="TOTAL"
    ).font = fonte_total

    # VALOR
    ws.cell(
        row=linha_total_notas,
        column=4,
        value=f"=SUM(D{linha_inicio_notas}:D{linha_fim_notas})"
    )

    # IR
    ws.cell(
        row=linha_total_notas,
        column=5,
        value=f"=SUM(E{linha_inicio_notas}:E{linha_fim_notas})"
    )

    # INSS
    ws.cell(
        row=linha_total_notas,
        column=6,
        value=f"=SUM(F{linha_inicio_notas}:F{linha_fim_notas})"
    )

    # PIS / COFINS / CSLL
    ws.cell(
        row=linha_total_notas,
        column=7,
        value=f"=SUM(G{linha_inicio_notas}:G{linha_fim_notas})"
    )

    # VALOR LÍQUIDO
    ws.cell(
        row=linha_total_notas,
        column=8,
        value=f"=SUM(H{linha_inicio_notas}:H{linha_fim_notas})"
    )

    # FORMATAÇÃO
    for coluna in range(3, 9):

        celula = ws.cell(
            row=linha_total_notas,
            column=coluna
        )

        celula.font = fonte_total
        celula.border = borda

        if coluna >= 4:
            celula.number_format = formato_moeda

    linha += 1
    # ========================================================
    # LARGURA DAS COLUNAS
    # ========================================================

    larguras = {
        "A": 14,
        "B": 40,
        "C": 14,
        "D": 16,
        "E": 14,
        "F": 14,
        "G": 20,
        "H": 18,
        "I": 30
    }

    for coluna, largura in larguras.items():

        ws.column_dimensions[coluna].width = largura

    # ========================================================
    # DEMONSTRATIVO
    # ========================================================

    linha += 2

    ws.cell(
        row=linha,
        column=1,
        value="DEMONSTRATIVO DE IMPOSTOS"
    ).font = fonte_titulo

    linha += 1

    ws.cell(
        row=linha,
        column=1,
        value="MÉDICO"
    ).font = fonte_cabecalho

    ws.cell(
        row=linha,
        column=2,
        value="VALOR DOS IMPOSTOS"
    ).font = fonte_cabecalho

    linha += 1

    linha_inicio_demonstrativo = linha

    # ========================================================
    # DEMONSTRATIVO POR MÉDICO
    # ========================================================

    # ========================================================
    # DEMONSTRATIVO POR MÉDICO
    # ========================================================

    linha_demonstrativo = linha_inicio_demonstrativo

    for medico in medicos:

        # Mesmo nome utilizado na criação da aba
        nome_aba = medico

        caracteres_invalidos = [
            "\\",
            "/",
            "*",
            "[",
            "]",
            ":",
            "?"
        ]

        for caractere in caracteres_invalidos:
            nome_aba = nome_aba.replace(
                caractere,
                ""
            )

        nome_aba = nome_aba[:31]

        # MÉDICO
        ws.cell(
            row=linha_demonstrativo,
            column=1,
            value=medico
        )

        # VALOR DOS IMPOSTOS
        # Será preenchido depois que a aba do médico
        # for criada e o TOTAL DOS IMPOSTOS estiver definido.
        ws.cell(
            row=linha_demonstrativo,
            column=2,
            value=None
        )

        ws.cell(
            row=linha_demonstrativo,
            column=2
        ).number_format = formato_moeda

        ws.cell(
            row=linha_demonstrativo,
            column=1
        ).border = borda

        ws.cell(
            row=linha_demonstrativo,
            column=2
        ).border = borda

        linha_demonstrativo += 1

    # ========================================================
    # TOTAL DO DEMONSTRATIVO
    # ========================================================

    linha_total_demonstrativo = linha_demonstrativo

    ws.cell(
        row=linha_total_demonstrativo,
        column=1,
        value="TOTAL"
    ).font = fonte_total

    ws.cell(
        row=linha_total_demonstrativo,
        column=2,
        value=(
            f"=SUM("
            f"B{linha_inicio_demonstrativo}:"
            f"B{linha_total_demonstrativo - 1}"
            f")"
        )
    )

    ws.cell(
        row=linha_total_demonstrativo,
        column=2
    ).font = fonte_total

    ws.cell(
        row=linha_total_demonstrativo,
        column=2
    ).number_format = formato_moeda

    # ========================================================
    # ABAS INDIVIDUAIS DOS MÉDICOS
    # ========================================================

    for medico in medicos:

        notas_medico = [
            nota
            for nota in dados_notas
            if nota.get("Medico", "").strip()
            == medico
        ]

        # ----------------------------------------------------
        # NOME DA ABA
        # ----------------------------------------------------

        nome_aba = medico

        caracteres_invalidos = [
            "\\",
            "/",
            "*",
            "[",
            "]",
            ":",
            "?"
        ]

        for caractere in caracteres_invalidos:
            nome_aba = nome_aba.replace(
                caractere,
                ""
            )

        nome_aba = nome_aba[:31]

        ws_medico = wb.create_sheet(
            title=nome_aba
        )

        # ----------------------------------------------------
        # TÍTULO
        # ----------------------------------------------------

        ws_medico["A1"] = medico
        ws_medico["A1"].font = fonte_titulo

        ws_medico["A2"] = "FATURAMENTO - AGOSTO - 2026"
        ws_medico["A2"].font = fonte_cabecalho

        # ----------------------------------------------------
        # CABEÇALHO FATURAMENTO
        # ----------------------------------------------------

        linha_medico = 4

        cabecalhos_medico = [
            "DATA",
            "CLIENTE",
            "N.FISCAL",
            "VALOR",
            "IR",
            "PIS-COFINS",
            "TAXA ADM",
            "VALOR LÍQ."
        ]

        for coluna, cabecalho in enumerate(
            cabecalhos_medico,
            start=1
        ):

            celula = ws_medico.cell(
                row=linha_medico,
                column=coluna,
                value=cabecalho
            )

            celula.font = fonte_cabecalho
            celula.alignment = alinhamento_centro
            celula.border = borda

        linha_medico += 1

        linha_inicio_faturamento = linha_medico

        # ----------------------------------------------------
        # NOTAS DO MÉDICO
        # ----------------------------------------------------

        for nota in notas_medico:

            situacao_nota = str(
                nota.get("SituacaoNota", "")
            ).strip().upper()

            if situacao_nota == "C":
                continue

            valor = float(
                nota.get("ValorTotal", 0)
            )

            ir = float(
                nota.get("ValorIr", 0)
            )

            pis = float(
                nota.get("ValorPis", 0)
            )

            cofins = float(
                nota.get("ValorCofins", 0)
            )

            csll = float(
                nota.get("ValorCsll", 0)
            )

            situacao_nota = str(
                nota.get("SituacaoNota", "")
            ).strip().upper()

            valor_pcc_retido = nota.get(
                "ValorPccRetido"
            )

            if situacao_nota == "C":

                pis_cofins = 0

            elif (
                valor_pcc_retido is not None
                and valor_pcc_retido != ""
            ):

                pis_cofins = float(
                    valor_pcc_retido
                )

            elif csll > 0:

                pis_cofins = round(
                    valor * 0.0465,
                    2
                )

            else:

                pis_cofins = 0

            taxa_adm = 0.00

            valor_liquido = (
                valor
                - ir
                - pis_cofins
                - taxa_adm
            )

            valores = [
                nota.get("DataEmissao", ""),
                nota.get("TomadorServico", ""),
                nota.get("NumeroNF", ""),
                valor,
                ir,
                pis_cofins,
                taxa_adm,
                valor_liquido
            ]

            for coluna, valor_celula in enumerate(
                valores,
                start=1
            ):

                celula = ws_medico.cell(
                    row=linha_medico,
                    column=coluna,
                    value=valor_celula
                )

                celula.border = borda

                if coluna >= 4:

                    celula.number_format = formato_moeda

            linha_medico += 1

        linha_fim_faturamento = (
            linha_medico - 1
        )

        # ----------------------------------------------------
        # TOTAIS PARA O RESUMO
        # ----------------------------------------------------

        notas_validas = [
            nota
            for nota in notas_medico
            if str(
                nota.get("SituacaoNota", "")
            ).strip().upper() != "C"
        ]

        total_valor = sum(
            float(nota.get("ValorTotal", 0))
            for nota in notas_validas
        )

        total_ir = sum(
            float(nota.get("ValorIr", 0))
            for nota in notas_validas
        )

        total_pis_cofins = 0

        for nota in notas_validas:

            valor_pcc_retido = nota.get(
                "ValorPccRetido"
            )

            csll = float(
                nota.get("ValorCsll", 0)
            )

            if (
                valor_pcc_retido is not None
                and valor_pcc_retido != ""
            ):

                pcc = float(
                    valor_pcc_retido
                )

            elif csll > 0:

                pcc = round(
                    float(nota.get("ValorTotal", 0))
                    * 0.0465,
                    2
                )

            else:

                pcc = 0

            total_pis_cofins += pcc

        total_taxa_adm = 0.00

        total_liquido = (
            total_valor
            - total_ir
            - total_pis_cofins
            - total_taxa_adm
        )
        # ----------------------------------------------------
        # ----------------------------------------------------
        # TOTAIS FATURAMENTO
        # ----------------------------------------------------

        linha_total_faturamento = linha_medico

        valores_totais = [
            "",
            "TOTAL",
            "",
            f"=SUM(D{linha_inicio_faturamento}:D{linha_fim_faturamento})",
            f"=SUM(E{linha_inicio_faturamento}:E{linha_fim_faturamento})",
            f"=SUM(F{linha_inicio_faturamento}:F{linha_fim_faturamento})",
            f"=SUM(G{linha_inicio_faturamento}:G{linha_fim_faturamento})",
            f"=SUM(H{linha_inicio_faturamento}:H{linha_fim_faturamento})"
        ]

        for coluna, valor_celula in enumerate(
            valores_totais,
            start=1
        ):

            celula = ws_medico.cell(
                row=linha_total_faturamento,
                column=coluna,
                value=valor_celula
            )

            celula.font = fonte_total
            celula.border = borda

            if coluna >= 4:
                celula.number_format = formato_moeda

        # ----------------------------------------------------
        # IMPOSTOS
        # ----------------------------------------------------

        linha_medico += 2

        ws_medico.cell(
            row=linha_medico,
            column=1,
            value="IMPOSTOS"
        ).font = fonte_titulo

        linha_medico += 1

        # ----------------------------------------------------
        # VALORES DESTACADOS
        # ----------------------------------------------------

        total_iss_destacado = sum(
            float(nota.get("ValorIss", 0))
            for nota in notas_validas
        )

        total_pis_destacado = sum(
            float(nota.get("ValorPis", 0))
            for nota in notas_validas
        )

        total_cofins_destacado = sum(
            float(nota.get("ValorCofins", 0))
            for nota in notas_validas
        )

        total_csll_destacado = sum(
            float(nota.get("ValorCsll", 0))
            for nota in notas_validas
        )

        total_ir_destacado = sum(
            float(nota.get("ValorIr", 0))
            for nota in notas_validas
        )

        # ----------------------------------------------------
        # CÁLCULOS
        # ----------------------------------------------------

        base_calculo = total_valor

        base_formatada = (
            f"{base_calculo:,.2f}"
            .replace(",", "X")
            .replace(".", ",")
            .replace("X", ".")
        )

        # ----------------------------------------------------
        # NÚMEROS DAS NOTAS
        # ----------------------------------------------------

        numero_inicial = (
            notas_medico[0].get("NumeroNF", "")
            if notas_medico
            else ""
        )

        numero_final = (
            notas_medico[-1].get("NumeroNF", "")
            if notas_medico
            else ""
        )

        base_formatada = (
            f"{base_calculo:,.2f}"
            .replace(",", "X")
            .replace(".", ",")
            .replace("X", ".")
        )

        # ----------------------------------------------------
        # TABELA DE IMPOSTOS
        # ----------------------------------------------------

        cabecalhos_impostos = [
            "COD.REC.",
            "IMPOSTOS",
            "VENCTO",
            "VALOR IMPOSTO",
            "DEDUÇÃO NF",
            "TOTAL"
        ]

        for coluna, cabecalho in enumerate(
            cabecalhos_impostos,
            start=1
        ):

            celula = ws_medico.cell(
                row=linha_medico,
                column=coluna,
                value=cabecalho
            )

            celula.font = fonte_cabecalho
            celula.alignment = alinhamento_centro
            celula.border = borda

        linha_medico += 1

        dados_impostos = [

            [
                "",
                "HONORÁRIO - 08/2026",
                "25/09/2026",
                0,
                0,
                honorario_por_medico
            ],

            [
                "04030",
                (
                    "ISS - SP - 2% - "
                    f"NF.{numero_inicial} A {numero_final} - "
                    f"Base de Calc R$ {base_formatada}"
                ),
                "10/09/2026",
                None,
                0,
                None
            ],

            [
                "",
                "DARF GPS - 08/2026",
                "18/09/2026",
                0,
                0,
                darf_gps_por_medico
            ],

            [
                "2172",
                (
                    "COFINS - 3% - "
                    f"Base de Cálculo R$ {base_formatada}"
                ),
                "25/09/2026",
                None,
                None,
                None
            ],

            [
                "8109",
                (
                    "PIS - 0,65% - "
                    f"Base de Cálculo R$ {base_formatada}"
                ),
                "25/09/2026",
                None,
                None,
                None
            ],

            [
                "2372",
                (
                    "CONTRIBUIÇÃO SOCIAL - "
                    "Alíq.Reduzida - 1,08% - "
                    f"Base de Cálculo R$ {base_formatada}"
                ),
                "30/09/2026",
                None,
                None,
                None
            ],

            [
                "2089",
                (
                    "I.R.P.J. - "
                    "Alíq.Reduzida - 1,2% - "
                    f"Base de Cálculo R$ {base_formatada}"
                ),
                "30/09/2026",
                None,
                None,
                0
            ]
        ]

        linha_inicio_impostos = linha_medico

        for dados in dados_impostos:

            for coluna, valor_celula in enumerate(
                dados,
                start=1
            ):

                celula = ws_medico.cell(
                    row=linha_medico,
                    column=coluna,
                    value=valor_celula
                )

                celula.border = borda

                if coluna >= 4:
                    celula.number_format = formato_moeda

            linha_medico += 1


        linha_honorario = linha_inicio_impostos
        linha_iss = linha_inicio_impostos + 1
        linha_gps = linha_inicio_impostos + 2
        linha_cofins = linha_inicio_impostos + 3
        linha_pis = linha_inicio_impostos + 4
        linha_csll = linha_inicio_impostos + 5
        linha_ir = linha_inicio_impostos + 6


        # ISS
        ws_medico.cell(
            row=linha_iss,
            column=4,
            value=f"=D{linha_total_faturamento}*2%"
        )

        ws_medico.cell(
            row=linha_iss,
            column=6,
            value=f"=D{linha_iss}-E{linha_iss}"
        )


        # GPS
        ws_medico.cell(
            row=linha_gps,
            column=6,
            value=f"=D{linha_gps}-E{linha_gps}"
        )


        # COFINS
        ws_medico.cell(
            row=linha_cofins,
            column=4,
            value=f"=D{linha_total_faturamento}*3%"
        )

        ws_medico.cell(
            row=linha_cofins,
            column=5,
            value=(
                f"=(100%*F{linha_total_faturamento}/4.65%)*3%"
            )
        )

        ws_medico.cell(
            row=linha_cofins,
            column=6,
            value=f"=D{linha_cofins}-E{linha_cofins}"
        )


        # PIS
        ws_medico.cell(
            row=linha_pis,
            column=4,
            value=f"=D{linha_total_faturamento}*0.65%"
        )

        ws_medico.cell(
            row=linha_pis,
            column=5,
            value=(
                f"=(100%*F{linha_total_faturamento}/4.65%)*0.65%"
            )
        )

        ws_medico.cell(
            row=linha_pis,
            column=6,
            value=f"=D{linha_pis}-E{linha_pis}"
        )


        # CSLL
        ws_medico.cell(
            row=linha_csll,
            column=4,
            value=f"=D{linha_total_faturamento}*1.08%"
        )

        ws_medico.cell(
            row=linha_csll,
            column=5,
            value=(
                f"=(100%*F{linha_total_faturamento}/4.65%)*1%"
            )
        )

        ws_medico.cell(
            row=linha_csll,
            column=6,
            value=f"=D{linha_csll}-E{linha_csll}"
        )


        # IR
        ws_medico.cell(
            row=linha_ir,
            column=4,
            value=f"=D{linha_total_faturamento}*1.2%"
        )

        ws_medico.cell(
            row=linha_ir,
            column=5,
            value=f"=E{linha_total_faturamento}"
        )

        ws_medico.cell(
            row=linha_ir,
            column=6,
            value=f"=MAX(0,D{linha_ir}-E{linha_ir})"
        )

        # ----------------------------------------------------
        # TOTAL DOS IMPOSTOS
        # ----------------------------------------------------

        linha_total_impostos = linha_medico

        ws_medico.cell(
            row=linha_total_impostos,
            column=1,
            value="TOTAL"
        ).font = fonte_total

        celula = ws_medico.cell(
            row=linha_total_impostos,
            column=6,
            value=(
                f"=SUM("
                f"F{linha_inicio_impostos}:"
                f"F{linha_total_impostos - 1}"
                f")"
            )
        )

        celula.font = fonte_total
        celula.number_format = formato_moeda

        # ----------------------------------------------------
        # LIGA TOTAL DA ABA AO DEMONSTRATIVO
        # ----------------------------------------------------

        for linha_demo in range(
            linha_inicio_demonstrativo,
            linha_total_demonstrativo
        ):

            if ws.cell(
                row=linha_demo,
                column=1
            ).value == medico:

                ws.cell(
                    row=linha_demo,
                    column=2,
                    value=(
                        f"='{nome_aba}'!"
                        f"F{linha_total_impostos}"
                    )
                )

                ws.cell(
                    row=linha_demo,
                    column=2
                ).number_format = formato_moeda

                break
            
        # ----------------------------------------------------
        # LARGURA DAS COLUNAS
        # ----------------------------------------------------

        larguras_medico = {
            "A": 14,
            "B": 45,
            "C": 15,
            "D": 18,
            "E": 18,
            "F": 18,
            "G": 18,
            "H": 18
        }

        for coluna, largura in larguras_medico.items():

            ws_medico.column_dimensions[
                coluna
            ].width = largura

    # ========================================================
    # SALVA
    # ========================================================

    wb.save(caminho_saida)

    return caminho_saida